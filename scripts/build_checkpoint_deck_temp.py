from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT = "submissions/2026-10-29-checkpoint-deck.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(15, 23, 42)
BLUE = RGBColor(37, 99, 235)
CYAN = RGBColor(6, 182, 212)
GREEN = RGBColor(22, 163, 74)
AMBER = RGBColor(217, 119, 6)
RED = RGBColor(220, 38, 38)
SLATE = RGBColor(71, 85, 105)
LIGHT = RGBColor(241, 245, 249)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(203, 213, 225)

def bg(slide, color=NAVY):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid(); shape.fill.fore_color.rgb = color
    shape.line.fill.background()

def txt(slide, text, x, y, w, h, size=24, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box

def title(slide, kicker, heading, sub=None):
    txt(slide, kicker.upper(), 0.7, 0.42, 4.8, 0.35, 11, True, CYAN)
    txt(slide, heading, 0.7, 0.85, 11.8, 0.8, 28, True)
    if sub:
        txt(slide, sub, 0.72, 1.62, 11.7, 0.5, 14, False, MUTED)

def card(slide, x, y, w, h, header, body, accent=BLUE):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor(30, 41, 59)
    s.line.color.rgb = RGBColor(51, 65, 85)
    a = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = accent; a.line.fill.background()
    txt(slide, header, x+0.25, y+0.22, w-0.45, 0.36, 15, True)
    txt(slide, body, x+0.25, y+0.68, w-0.45, h-0.8, 12, False, MUTED)

def chip(slide, label, x, y, color=GREEN, w=1.4):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.38))
    s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()
    txt(slide, label, x, y+0.03, w, 0.25, 10, True, WHITE, PP_ALIGN.CENTER)

def evidence(slide, x, y, w, h, header, lines, accent=CYAN):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor(2, 6, 23)
    s.line.color.rgb = accent
    txt(slide, header, x+0.25, y+0.18, w-0.5, 0.35, 13, True, accent)
    yy = y+0.68
    for line in lines:
        txt(slide, line, x+0.28, yy, w-0.55, 0.34, 11, False, RGBColor(226,232,240))
        yy += 0.43

def add_footer(slide, n):
    txt(slide, f"PROOFSHAPE  •  CHECKPOINT  •  {n}/10", 0.7, 7.05, 4.5, 0.22, 9, True, RGBColor(148,163,184))

# 1
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
txt(s,"PROOFSHAPE",0.8,0.65,5.5,0.65,18,True,CYAN)
txt(s,"Smartphone 3D Inspection\nCheckpoint Review",0.8,1.5,8.4,1.6,34,True)
txt(s,"Mobile capture → metric reconstruction → provenance-aware inspection",0.82,3.25,8.8,0.5,16,False,MUTED)
chip(s,"13-WEEK CAPSTONE",0.82,4.25,BLUE,1.75)
chip(s,"3-PERSON TEAM",2.78,4.25,GREEN,1.65)
chip(s,"OCT 29, 2026",4.62,4.25,AMBER,1.55)
card(s,9.35,1.45,3.1,3.7,"Current checkpoint","Phone capture is connected to the real service. Metric alignment, confidence-gated provenance, uncertainty, and AI frame supervision are working. CAD registration is the next inspection milestone.",CYAN)
add_footer(s,1)

# 2
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,"team","Three contributors, one inspection pipeline")
card(s,0.8,2.2,3.75,3.9,"Mamadou Jallow  ·  @mbj1994","Reconstruction & integration\n• Provenance / uncertainty\n• Metric alignment support\n• VLM supervision\n• CAD loading + I-03 registration",BLUE)
card(s,4.8,2.2,3.75,3.9,"Yashi Dalwala  ·  @dalwalyk","Capture & inspection\n• Capture → service integration\n• CAD / OBJ loading\n• Device validation\n• Cross-machine testing",GREEN)
card(s,8.8,2.2,3.75,3.9,"Tabeen Raoof  ·  @TabeenRaoof","Reconstruction validation\n• Golden-set reruns\n• Metric / board validation\n• Review & clean checkout\n• Measurement evidence",AMBER)
add_footer(s,2)

# 3
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,"progress 1/4","Phone capture now reaches the real reconstruction service","C-13 closed the gap between the browser capture flow and recon.service.")
evidence(s,0.8,2.2,6.0,3.9,"REAL-PHONE EVIDENCE",[
"iPhone → HTTPS dev proxy → recon.service",
"14-frame clean-checkout session completed on M3 Pro",
"metric: true  •  ChArUco reference tier",
"observed_fraction: 0.1395",
"638 KB GLB produced on disk"
],GREEN)
card(s,7.1,2.2,5.3,1.7,"What changed","EXIF-less phone frames are calibrated from the printed ChArUco board; the same-origin proxy lets the phone reach the real service without a second certificate.",CYAN)
card(s,7.1,4.15,5.3,1.95,"What we learned","The end-to-end path works, but phone reliability is not solved yet: early tests exposed board-calibration and fusion failure modes. Failures now report their real cause instead of failing silently.",AMBER)
add_footer(s,3)

# 4
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,"progress 2/4","Metric alignment is now board-anchored","R-23 removed the old camera-centre scale bias across all five golden objects.")
evidence(s,0.8,2.2,7.2,3.9,"R-23 GOLDEN-SET SCALE",[
"s-01   board scale 0.9980   RMS 2.540 mm",
"s-02   board scale 0.9971   RMS 3.016 mm",
"s-03   board scale 0.9994   RMS 1.615 mm",
"s-04   board scale 0.9987   RMS 2.084 mm",
"s-05   board scale 0.9982   RMS 2.356 mm"
],BLUE)
card(s,8.3,2.2,4.1,1.75,"Before","Legacy camera-centre scaling was roughly 0.907–0.958 on the golden set, making reconstructions visibly small.",RED)
card(s,8.3,4.2,4.1,1.9,"After","Board-anchored scale is roughly 0.997–0.999. This gives provenance and future CAD inspection a defensible metric frame.",GREEN)
add_footer(s,4)

# 5
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,"progress 3/4","Provenance is becoming stricter—not more optimistic","R-24 and R-25 protect the hard rule: unseen geometry can never pass or fail a part.")
evidence(s,0.8,2.2,5.9,3.95,"R-24 CONFIDENCE FILTER",[
"s-01  observations −3.1%",
"s-02  observations −4.0%",
"s-03  observations −4.5%",
"s-04  observations −4.2%",
"s-05  observations −2.8%",
"Only trusted reconstruction pixels may vouch for a vertex."
],CYAN)
card(s,7.0,2.2,5.4,1.65,"R-25 facing test","A vertex must now face the observing camera and agree with measured depth. Open-mesh normal direction uses only trusted, depth-consistent visible pairs.",BLUE)
card(s,7.0,4.1,5.4,2.05,"Final coverage semantics","observed_fraction is computed after the sigma downgrade, weighted by triangle area, from the same final labels encoded in the GLB served to the user.",GREEN)
add_footer(s,5)

# 6
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,"progress 4/4","Generative AI supervises sampled capture frames","A-01 added real Gemini API checks without letting the model decide geometry.")
card(s,0.8,2.2,4.0,3.9,"What the VLM does","• Samples selected capture frames\n• Flags wrong-part / screen / obvious scene issues\n• Validates structured responses\n• Stores per-frame + aggregate evidence\n• Fails safely when unavailable",BLUE)
evidence(s,5.1,2.2,3.4,3.9,"REAL API CHECK",[
"30 JPEG golden frames available",
"11 valid Gemini replies recorded",
"screen / different-part controls checked",
"token usage recorded",
"evidence JSON persisted"
],GREEN)
card(s,8.8,2.2,3.6,3.9,"Safety boundary","The VLM can contribute evidence and later veto an unsafe Pass, but it never computes metric geometry. Geometry, provenance and tolerances stay deterministic.",AMBER)
add_footer(s,6)

# 7
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,"timeline","From capture prototype to inspection pipeline")
items=[
("Weeks 1–3","Capture + dataset","Phone permissions, guided capture, golden objects",GREEN),
("Weeks 4–6","Reconstruction","VGGT, board pose, metric alignment, TSDF",BLUE),
("Weeks 7–9","Evidence","Cleanup, provenance, uncertainty, endpoint",CYAN),
("Weeks 10–11","Integration","Phone → real service, AI frame supervision",AMBER),
("Weeks 12–13","Inspection","CAD registration, ICP, verdict/report close-out",RED),
]
y=2.15
for i,(a,b,c,col) in enumerate(items):
    chip(s,a,0.8,y,col,1.55)
    txt(s,b,2.65,y-0.02,2.2,0.3,14,True)
    txt(s,c,4.75,y-0.02,7.3,0.4,12,False,MUTED)
    if i<len(items)-1:
        ln=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(1.52),Inches(y+0.42),Inches(0.035),Inches(0.5))
        ln.fill.solid(); ln.fill.fore_color.rgb=RGBColor(71,85,105); ln.line.fill.background()
    y += 0.92
add_footer(s,7)

# 8
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,"what's next 1/2","Close R-25 with measured evidence—not assumptions")
card(s,0.8,2.15,3.7,3.9,"1  Golden rerun","Recompute R-23 board alignment, re-fuse once, then compare current-main vs R-25 on the exact same cleaned mesh for s-01…s-05.",BLUE)
card(s,4.82,2.15,3.7,3.9,"2  Record the effect","Observed vertices, area-weighted final coverage, depth-gate rejection counts, sigma downgrades, and historical R-24 context.",CYAN)
card(s,8.84,2.15,3.7,3.9,"3  Final close-out","Add measured hours, mark R-25 Done, run final CI, get teammate approval on the final head, then merge.",GREEN)
add_footer(s,8)

# 9
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,"what's next 2/2","Move from reconstruction into CAD inspection")
card(s,0.8,2.15,3.75,3.9,"I-03  Global registration","Generate candidate CAD poses, keep metric scale fixed, and report free-scale residual as an integrity signal.",BLUE)
card(s,4.8,2.15,3.75,3.9,"I-04 / I-05  Refine","Robust point-to-plane ICP with Tukey loss, then explicitly detect orientation ambiguity instead of hiding it.",CYAN)
card(s,8.8,2.15,3.75,3.9,"End goal","Observed geometry → CAD deviation → Pass / Fail / Rescan / Unverifiable, with evidence that explains exactly why.",GREEN)
add_footer(s,9)

# 10
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s); title(s,"references","Project evidence used for this checkpoint")
refs=[
"ProofShape GitHub repository — stories, decisions, tests, pull requests and CI",
"R-23 — board-anchored metric alignment golden-set evidence",
"R-24 — confidence-filtered provenance golden-set evidence",
"C-13 / PR #146 — real-phone → real-service integration evidence",
"A-01 / PR #207 — sampled Gemini VLM supervision evidence",
"R-25 / PR #213 — facing, sigma truncation and final observed coverage",
"ProofShape proposal, manual and frozen OpenAPI contract"
]
y=2.05
for i,r in enumerate(refs,1):
    txt(s,f"{i:02d}",0.9,y,0.55,0.3,12,True,CYAN)
    txt(s,r,1.6,y-0.02,10.7,0.42,13,False,RGBColor(226,232,240))
    y+=0.58
txt(s,"Checkpoint deck • 10 slides • designed for a 5–7 minute presentation",0.9,6.35,10.5,0.4,13,True,MUTED)
add_footer(s,10)

prs.save(OUT)
print(OUT)
